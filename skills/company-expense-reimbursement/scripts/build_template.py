#!/usr/bin/env python
"""Build the reusable Excel workbook for the company paper reimbursement form."""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Mapping

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins

from reimbursement_profile import load_company_profile, load_user_profile

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "assets" / "company-expense-template.xlsx"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE_YELLOW = "FFF2CC"
PALE_RED = "FCE4D6"
WHITE = "FFFFFF"
GRAY = "666666"
thin = Side(style="thin", color="7F8C8D")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def build_reimbursement_page(ws, user: Mapping[str, str], company: Mapping[str, str], page_number: int = 1) -> None:
    """Create one fixed-size page matching the supplied paper form."""
    title_font = Font(name="Microsoft YaHei", size=20, bold=True, color=BLUE)
    label_font = Font(name="Microsoft YaHei", size=11, bold=True, color=BLUE)
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
    body_font = Font(name="Microsoft YaHei", size=10, color="000000")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    blue_fill = PatternFill("solid", fgColor=BLUE)
    yellow_fill = PatternFill("solid", fgColor=PALE_YELLOW)

    ws.title = "报销单" if page_number == 1 else f"报销单-第{page_number}页"
    ws.sheet_view.showGridLines = False
    # Four narrow date columns and wider label/value columns follow the paper form.
    # Keep the four date/time cells narrow, give the location field enough
    # room, and keep the final note field close to the supplied paper form.
    widths = [9, 9, 9, 9, 18, 18, 12, 12, 11, 11, 11, 11, 11, 6, 6, 6]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + index)].width = width
    for row in range(1, 19):
        ws.row_dimensions[row].height = 25
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[5].height = 31
    ws.row_dimensions[7].height = 29
    ws.row_dimensions[8].height = 29

    merges = [
        "A1:C2", "D1:L2", "M1:N1", "O1:P1", "M2:N2", "O2:P2",
        "A3:B3", "C3:F3", "G3:H3", "I3:J3", "K3:L3", "M3:N3", "O3:P3",
        "A4:B4", "C4:F4", "G4:H4", "I4:J4", "K4:L4", "M4:P4",
        "A5:B5", "C5:D5", "E5:F5", "G5:H5", "I5:J5", "K5:L5", "M5:M5", "N5:O5", "P5:P5",
        "A7:D7", "E7:F8", "G7:H7", "I7:I8", "J7:J8", "K7:K8", "L7:M7", "N7:P8",
        "A15:F15", "A16:H16", "I16:N16", "A18:C18", "G18:I18", "M18:P18",
    ]
    for merge in merges:
        ws.merge_cells(merge)

    ws["A1"] = ""
    ws["D1"] = "差旅费报销单"
    ws["D1"].font = title_font
    ws["D1"].alignment = center
    labels = {
        "M1": "预领款", "M2": "补领或缴还", "A3": "部门", "G3": "年", "I3": "月", "K3": "日",
        "A4": "姓名", "G4": "职别", "K4": "出差事由", "A5": "出差起止日期", "C5": "自", "G5": "至",
        "K5": "共", "M5": "天", "N5": "附单据", "A7": "起讫", "E7": "起讫地点", "G7": "车船费",
        "G8": "名称", "H8": "金额", "I7": "住宿费", "J7": "出差补助", "K7": "市内交通费",
        "L7": "杂费", "L8": "用途", "M8": "金额", "N7": "附注", "A15": "合计",
        "A16": "合计金额（大写）", "A18": "单位主管：", "G18": "复核：", "M18": "出差人：",
    }
    for addr, value in labels.items():
        ws[addr] = value
        ws[addr].font = label_font
        ws[addr].alignment = center

    # Every blue table header is white, including the nested 起讫 row.
    for row in (7, 8):
        for col in range(1, 17):
            ws.cell(row, col).fill = blue_fill
            ws.cell(row, col).font = header_font
            ws.cell(row, col).alignment = center
    for addr, value in {"A8": "日", "B8": "时", "C8": "日", "D8": "时"}.items():
        ws[addr] = value
        ws[addr].fill = blue_fill
        ws[addr].font = header_font
        ws[addr].alignment = center

    ws["C3"] = user.get("department", "")
    ws["C4"] = user.get("name", "")
    ws["I4"] = user.get("title", "")
    for addr in ("C3", "C4", "I4", "M4", "E5", "I5", "L5", "O1", "O2", "O3", "P5"):
        ws[addr].fill = yellow_fill
        ws[addr].font = body_font
        ws[addr].alignment = left

    for row in range(9, 15):
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
        ws.merge_cells(start_row=row, start_column=14, end_row=row, end_column=16)
        for col in range(1, 17):
            cell = ws.cell(row, col)
            cell.border = border
            cell.font = body_font
            cell.alignment = left if col in (5, 6, 14, 15, 16) else center
            if col not in (8, 9, 10, 11, 13):
                cell.fill = yellow_fill
            if col in (8, 9, 10, 11, 13):
                cell.number_format = "#,##0.00"

    for row in range(1, 19):
        for col in range(1, 17):
            ws.cell(row, col).border = border
    for col, letter in ((8, "H"), (9, "I"), (11, "K"), (13, "M")):
        ws.cell(15, col).value = f"=SUM({letter}9:{letter}14)"
        ws.cell(15, col).number_format = "#,##0.00"
        ws.cell(15, col).font = label_font
        ws.cell(15, col).alignment = center
    ws["A15"].font = label_font
    ws["A15"].alignment = center
    ws["A16"].font = label_font
    ws["A16"].alignment = left
    ws["O16"] = "￥"
    ws["O16"].alignment = center
    for addr in ("A18", "G18", "M18"):
        ws[addr].font = body_font
        ws[addr].alignment = left

    ws.freeze_panes = "A9"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.2, right=0.2, top=0.25, bottom=0.25, header=0.1, footer=0.1)
    ws.print_area = "A1:P18"


def add_reimbursement_page(wb: Workbook, user: Mapping[str, str], company: Mapping[str, str], page_number: int):
    ws = wb.create_sheet()
    build_reimbursement_page(ws, user, company, page_number)
    return ws


def make_workbook(company: Mapping[str, str] | None = None, user: Mapping[str, str] | None = None) -> Workbook:
    company = dict(company or {})
    user = dict(user or {})
    wb = Workbook()
    ws = wb.active
    build_reimbursement_page(ws, user, company)
    detail = wb.create_sheet("发票明细")
    check = wb.create_sheet("校验与说明")
    title_font = Font(name="Microsoft YaHei", size=18, bold=True, color=BLUE)
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
    body_font = Font(name="Microsoft YaHei", size=10, color="000000")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    blue_fill = PatternFill("solid", fgColor=BLUE)
    yellow_fill = PatternFill("solid", fgColor=PALE_YELLOW)
    red_fill = PatternFill("solid", fgColor=PALE_RED)

    headers = ["来源文件", "页码", "单据类型", "发票号码/订单号", "开票/乘车/入住日期", "结束日期", "购买方名称", "购买方税号", "购买方地址", "购买方电话", "购买方开户行", "购买方账号", "销方/承运方/酒店", "税率", "不含税金额", "税额", "价税合计", "出发地", "到达地", "入住时间", "离开时间", "分类", "可报销", "校验状态", "备注"]
    for col, value in enumerate(headers, 1):
        cell = detail.cell(1, col, value)
        cell.fill = blue_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border
    for row in range(2, 102):
        for col in range(1, 26):
            cell = detail.cell(row, col)
            cell.font = body_font
            cell.border = border
            cell.alignment = left if col in (1, 4, 7, 9, 13, 18, 19, 24, 25) else center
            if col in (15, 16, 17):
                cell.number_format = "#,##0.00"
    for col, width in {1: 34, 2: 8, 3: 16, 4: 24, 5: 18, 6: 16, 7: 24, 8: 22, 9: 42, 10: 14, 11: 18, 12: 24, 13: 30, 14: 10, 15: 14, 16: 14, 17: 14, 18: 18, 19: 18, 20: 16, 21: 16, 22: 14, 23: 12, 24: 20, 25: 40}.items():
        detail.column_dimensions[chr(64 + col)].width = width
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = "A1:Y101"
    category = DataValidation(type="list", formula1='"出租车/网约车,飞机/行程单,住宿,火车/动车,其他"', allow_blank=True)
    yes = DataValidation(type="list", formula1='"是,否,待确认"', allow_blank=True)
    status = DataValidation(type="list", formula1='"已核对,购买方信息不一致,待补充入住/离店时间,路线未闭环,金额待确认,来源无法读取"', allow_blank=True)
    for validation, ref in ((category, "V2:V101"), (yes, "W2:W101"), (status, "X2:X101")):
        detail.add_data_validation(validation)
        validation.add(ref)
    detail.conditional_formatting.add("X2:X101", FormulaRule(formula=['AND($X2<>"",$X2<>"已核对")'], fill=red_fill))

    check.sheet_view.showGridLines = False
    check.merge_cells("A1:D1")
    check["A1"] = "报销校验与填写说明"
    check["A1"].font = title_font
    check["A1"].alignment = center
    for col, value in enumerate(["字段", "公司标准值", "本次票据值", "结果/备注"], 1):
        cell = check.cell(3, col, value)
        cell.fill = blue_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border
    company_fields = [("公司名称", company.get("name", "")), ("税号", company.get("tax_id", "")), ("地址", company.get("address", "")), ("电话", company.get("phone", "")), ("开户行", company.get("bank_name", "")), ("账号", company.get("bank_account", ""))]
    for row, (field, value) in enumerate(company_fields, 4):
        for col, item in enumerate((field, value, "", "待核对（本地配置）" if value else "待填写本地配置"), 1):
            cell = check.cell(row, col, item)
            cell.border = border
            cell.alignment = left
            cell.font = body_font
            if col in (2, 3):
                cell.fill = yellow_fill
    check.merge_cells("A12:D12")
    check["A12"] = "必须检查"
    check["A12"].fill = blue_fill
    check["A12"].font = header_font
    check["A12"].alignment = center
    rules = [
        "每个 PDF/图片先用本地 MinerU 解析，并在发票明细保留来源文件和页码。",
        "住宿票必须同时有入住和离店日期；缺任一日期就标记待补充，不得当作完整凭证。",
        "A 到 B 出差当天的住宿费写在 A 到 B 行；B 地多张住宿票按日期从上到下分行。",
        "路线必须闭环；缺少回程证据时标记路线未闭环，不虚构行程。",
        "报销单不打印签名，单位主管、复核、出差人栏留空，打印后手写。",
        "模板字段、填充坐标、纸面表头必须一致；第一条行程必须核对交通方式和去程方向。",
    ]
    for row, rule in enumerate(rules, 13):
        check.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        check.cell(row, 1, f"{row - 12}. {rule}")
        check.cell(row, 1).alignment = left
        check.cell(row, 1).font = body_font
        for col in range(1, 5):
            check.cell(row, col).border = border
    for col, width in {"A": 20, "B": 48, "C": 48, "D": 30}.items():
        check.column_dimensions[col].width = width
    check.freeze_panes = "A4"
    check.page_setup.orientation = "landscape"
    check.page_setup.fitToWidth = 1
    check.sheet_properties.pageSetUpPr.fitToPage = True
    return wb


def main() -> int:
    parser = argparse.ArgumentParser(description="Build a blank or locally populated reimbursement template.")
    parser.add_argument("--output", default=str(OUTPUT), help="Destination .xlsx path")
    parser.add_argument("--profile", help="Local company profile YAML")
    parser.add_argument("--user-profile", help="Local traveler profile YAML")
    args = parser.parse_args()
    output = Path(args.output).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    company, company_path = load_company_profile(args.profile, discover=bool(args.profile))
    user, user_path = load_user_profile(args.user_profile, discover=bool(args.user_profile))
    make_workbook(company, user).save(output)
    print(f"{output}\ncompany_profile={company_path or 'none'} user_profile={user_path or 'none'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
