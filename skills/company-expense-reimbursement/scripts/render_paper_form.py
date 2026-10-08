#!/usr/bin/env python
"""Render paper-form sheets from a reimbursement workbook as print-ready HTML."""

from __future__ import annotations

import argparse
import html
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from string import Template

from openpyxl import load_workbook

from reimbursement_amount import paper_total


# ── helpers ──────────────────────────────────────────────────────────────────

def display(value: object, *, money: bool = False) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return f"{value.month}月{value.day}日"
    if isinstance(value, date):
        return f"{value.month}月{value.day}日"
    if isinstance(value, time):
        return value.strftime("%H:%M" if value.second == 0 else "%H:%M:%S")
    if money and isinstance(value, (int, float, Decimal)):
        v = Decimal(str(value))
        return f"{v:,.2f}" if v else ""
    return str(value).strip()


def e(value: object, *, money: bool = False) -> str:
    return html.escape(display(value, money=money), quote=True)


def cv(ws, addr: str) -> object:
    return ws[addr].value


# ── expense row ───────────────────────────────────────────────────────────────

def expense_row(ws, row: int, last: bool = False) -> str:
    """Generate one expense row. If *last* is True, inject ¥: into the allowance cell."""
    date_parts: list[str] = []
    for date_cell, time_cell in ((f"A{row}", f"B{row}"), (f"C{row}", f"D{row}")):
        d = cv(ws, date_cell)
        t = cv(ws, time_cell)
        if d is not None or t is not None:
            parts = " ".join(p for p in (e(d), e(t)) if p)
            date_parts.append(parts)
    date_html = "<br>".join(date_parts)

    # allowance cell: blank (finance fills), but last row shows ¥：marker
    allowance_content = "¥：" if last else ""
    allowance_class = "ca yen-marker" if last else "ca"

    cols = [
        ("cd",  date_html,              False, True),   # 起讫日期（已转义）
        ("cl",  cv(ws, f"E{row}"),      False, False),  # 起讫地点
        ("cn",  cv(ws, f"G{row}"),      False, False),  # 车船名称
        ("cv",  cv(ws, f"H{row}"),      True,  False),  # 车船金额
        ("ch",  cv(ws, f"I{row}"),      True,  False),  # 宿费
        (None,  allowance_content,      False, True),   # 出差补助（财务填）
        ("co",  cv(ws, f"K{row}"),      True,  False),  # 市内交通
        ("cmn", cv(ws, f"L{row}"),      False, False),  # 杂费用途
        ("cma", cv(ws, f"M{row}"),      True,  False),  # 杂费金额
        ("cno", cv(ws, f"N{row}"),      False, False),  # 附注
    ]

    cells = ""
    for i, (css, val, money, raw) in enumerate(cols):
        if css is None:
            css = allowance_class
        if raw:
            cells += f'<td class="{css}">{val}</td>'
        else:
            cells += f'<td class="{css}" data-field="{css}-{row}">{e(val, money=money)}</td>'
    return f"<tr>{cells}</tr>"


# ── page ──────────────────────────────────────────────────────────────────────

def paper_page(ws) -> str:
    purpose = cv(ws, "M4") or cv(ws, "L4")

    def total(col: str) -> Decimal:
        return sum(
            Decimal(str(cv(ws, f"{col}{r}") or 0))
            for r in range(9, 15)
            if isinstance(cv(ws, f"{col}{r}"), (int, float, Decimal))
        )

    total_vehicle = total("H")
    total_hotel   = total("I")
    total_local   = total("K")
    total_misc    = total("M")

    report_date = "".join(display(cv(ws, a)) for a in ("G3", "I3", "K3"))
    if not report_date:
        report_date = "年　　月　　日"

    # rows 9-14: mark row 14 as last
    rows_html = "".join(
        expense_row(ws, r, last=(r == 14))
        for r in range(9, 15)
    )

    def money_td(css: str, val: Decimal) -> str:
        return f'<td class="{css}">{html.escape(display(val, money=True))}</td>'

    return f'''\
<section class="paper">
  <!-- 顶部：标题 + 日期 + 预领款 -->
  <div class="top">
    <div class="title-wrap">
      <h1>差 旅 费 报 销 单</h1>
      <div class="form-date">{e(report_date)}</div>
    </div>
    <div class="advance">
      <div class="adv-row">
        <span class="adv-label">预　领　款</span>
        <span class="adv-val" data-field="advance">{e(cv(ws, "O1"), money=True)}</span>
      </div>
      <div class="adv-row">
        <span class="adv-label">补款或缴还</span>
        <span class="adv-val" data-field="repay">{e(cv(ws, "O2"), money=True)}</span>
      </div>
    </div>
  </div>

  <!-- 部门 -->
  <div class="dept-row">
    <span>部门：</span>
    <span class="val" data-field="department">{e(cv(ws, "C3"))}</span>
  </div>

  <!-- 姓名 / 职别 / 出差事由 -->
  <div class="identity-row">
    <div class="lv">
      <span class="l">姓　名</span>
      <span class="v" data-field="name">{e(cv(ws, "C4"))}</span>
    </div>
    <div class="lv">
      <span class="l">职　别</span>
      <span class="v" data-field="role">{e(cv(ws, "I4"))}</span>
    </div>
    <div class="lv">
      <span class="l">出差&#10;事由</span>
      <span class="v" data-field="purpose">{e(purpose)}</span>
    </div>
  </div>

  <!-- 出差起止日期 / 天数 / 附单据 -->
  <div class="trip-row">
    <div class="lv">
      <span class="l">出差起&#10;止日期</span>
      <span class="v" data-field="trip-period">\
自　{e(cv(ws, "E5"))}　至　{e(cv(ws, "I5"))}</span>
    </div>
    <div class="days-cell">
      <span class="d0">共</span>
      <span class="d1" data-field="trip-days">{e(cv(ws, "K5"))}</span>
      <span class="d2">天</span>
    </div>
    <div class="lv">
      <span class="l">附单据</span>
      <span class="v" data-field="attachments">{e(cv(ws, "P5"))}&nbsp;张</span>
    </div>
  </div>

  <!-- 明细表 -->
  <table class="exp">
    <colgroup>
      <col class="cd"><col class="cl">
      <col class="cn"><col class="cv">
      <col class="ch"><col class="ca">
      <col class="co">
      <col class="cmn"><col class="cma">
      <col class="cno">
    </colgroup>
    <thead>
      <tr>
        <th rowspan="2">起讫<br>日期</th>
        <th rowspan="2">起讫地点</th>
        <th colspan="2">车　船　费</th>
        <th rowspan="2">宿　费</th>
        <th rowspan="2">出差<br>补助</th>
        <th rowspan="2">市内<br>交通费</th>
        <th colspan="2">杂　　费</th>
        <th rowspan="2">附　注</th>
      </tr>
      <tr><th>名称</th><th>金额</th><th>用途</th><th>金额</th></tr>
    </thead>
    <tbody>{rows_html}</tbody>
    <tfoot>
      <tr>
        <td colspan="3" class="tlabel">合　　计</td>
        {money_td("cv", total_vehicle)}
        {money_td("ch", total_hotel)}
        <td class="ca"></td>
        {money_td("co", total_local)}
        <td class="cmn"></td>
        {money_td("cma", total_misc)}
        <td class="cno"></td>
      </tr>
    </tfoot>
  </table>

  <!-- 合计金额大写 -->
  <div class="uc-row">
    <span class="uc-label">合计金额（大写）</span>
    <span class="uc-val" data-field="uppercase"></span>
  </div>

  <!-- 签名 -->
  <div class="sig-row">
    <div>单位主管：</div>
    <div>复　　核：</div>
    <div>出差人：</div>
  </div>
</section>'''


# ── main ─────────────────────────────────────────────────────────────────────

def main() -> int:
    parser = argparse.ArgumentParser(
        description="Render reimbursement paper sheets to standalone HTML."
    )
    parser.add_argument("workbook", help="Path to a reimbursement .xlsx workbook")
    parser.add_argument("--output", required=True, help="Destination HTML path")
    args = parser.parse_args()

    workbook_path = Path(args.workbook).expanduser().resolve()
    output = Path(args.output).expanduser().resolve()
    workbook = load_workbook(workbook_path, data_only=True)
    sheets = [s for s in workbook.worksheets if s.title.startswith("报销单")]
    if not sheets:
        raise ValueError("Workbook has no paper reimbursement sheet (title must start with '报销单')")

    template_path = (
        Path(__file__).resolve().parents[1] / "assets" / "reimbursement-paper-form.html"
    )
    template = Template(template_path.read_text(encoding="utf-8"))
    pages = "\n".join(paper_page(s) for s in sheets)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(template.substitute(pages=pages), encoding="utf-8")

    non_allowance_total = sum((paper_total(s) for s in sheets), Decimal("0"))
    print(f"{output}\npages={len(sheets)}\nnon_allowance_total={non_allowance_total:.2f}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
