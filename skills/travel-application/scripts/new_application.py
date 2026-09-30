#!/usr/bin/env python
"""Generate a filled company travel application from the local template.

The reusable template stays unchanged. Personal fields are loaded from the
private user profile unless explicit command-line values are provided.
"""

from __future__ import annotations

import argparse
import os
from datetime import date
from pathlib import Path
from typing import Any

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


USER_KEYS = ("name", "department", "title", "employee_id")


def read_yaml(path: Path) -> dict[str, Any]:
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    if not isinstance(data, dict):
        raise ValueError(f"YAML must contain a mapping: {path}")
    return data


def find_user_profile(explicit: str | None) -> Path | None:
    candidates: list[Path] = []
    if explicit:
        candidates.append(Path(explicit).expanduser())
    if os.environ.get("CODEX_USER_PROFILE"):
        candidates.append(Path(os.environ["CODEX_USER_PROFILE"]).expanduser())
    candidates.append(Path.home() / ".codex" / "private" / "user-profile.yaml")
    return next((p.resolve() for p in candidates if p.exists()), None)


def load_user_profile(explicit: str | None) -> tuple[dict[str, str], Path | None]:
    values = {key: "" for key in USER_KEYS}
    path = find_user_profile(explicit)
    if not path:
        return values, None
    data = read_yaml(path)
    source = data.get("user", data)
    aliases = {
        "name": ("name", "姓名"),
        "department": ("department", "部门"),
        "title": ("title", "职务", "职别"),
        "employee_id": ("employee_id", "工号"),
    }
    for key, names in aliases.items():
        for name in names:
            if source.get(name) not in (None, ""):
                values[key] = str(source[name])
                break
    return values, path


def load_company_name(explicit: str | None) -> str:
    candidates: list[Path] = []
    if explicit:
        candidates.append(Path(explicit).expanduser())
    if os.environ.get("CODEX_EXPENSE_PROFILE"):
        candidates.append(Path(os.environ["CODEX_EXPENSE_PROFILE"]).expanduser())
    candidates.append(Path.home() / ".codex" / "private" / "company-profile.yaml")
    for path in candidates:
        if path.exists():
            data = read_yaml(path)
            source = data.get("company", data)
            return str(source.get("name") or source.get("company_name") or "")
    return ""


def iso_date(value: str) -> date:
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise ValueError(f"日期必须使用 YYYY-MM-DD：{value}") from exc


def chinese_date(value: str) -> str:
    parsed = iso_date(value)
    return f"{parsed.year}年{parsed.month}月{parsed.day}日"


def date_range(start: str, end: str) -> str:
    start_d = iso_date(start)
    end_d = iso_date(end)
    if end_d < start_d:
        raise ValueError("结束日期不能早于开始日期")
    days = (end_d - start_d).days + 1
    return f"{chinese_date(start)}至{chinese_date(end)}，共{days}日"


def set_cell(cell, value: str, *, align=WD_ALIGN_PARAGRAPH.CENTER) -> None:
    cell.text = value
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        for run in paragraph.runs:
            run.font.name = "宋体"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(10.5)


def fill_form_table(table, args, *, route: str, trip_days: str) -> bool:
    if len(table.rows) != 10 or len(table.columns) != 6:
        return False

    unknown = "待补充"
    set_cell(table.cell(0, 1), args.name or unknown)
    set_cell(table.cell(0, 3), args.title or unknown)
    set_cell(table.cell(0, 5), args.department or unknown)
    set_cell(table.cell(1, 1), trip_days)
    set_cell(table.cell(2, 1), args.companion or unknown)

    set_cell(table.cell(4, 1), f"{chinese_date(args.start)}至{chinese_date(args.end)}")
    set_cell(table.cell(4, 2), args.destination)
    reason = f"路线：{route}；{args.purpose}"
    set_cell(table.cell(4, 3), reason, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell(table.cell(4, 5), args.transport)
    for col in (1, 2, 3, 5):
        set_cell(table.cell(5, col), "")
    for col in range(1, 6):
        set_cell(table.cell(7, col), getattr(args, f"budget_{col}") or unknown)

    # Approval and signature cells remain blank for manual/external approval.
    for col in (1, 3, 5):
        set_cell(table.cell(8, col), "")

    if args.actual_start and args.actual_end:
        set_cell(table.cell(9, 1), f"{chinese_date(args.actual_start)}至{chinese_date(args.actual_end)}")
    else:
        set_cell(table.cell(9, 1), "")
    set_cell(table.cell(9, 5), "")
    return True


def update_case(case_dir: Path, *, output: Path, args, route: str) -> None:
    case_path = case_dir / "case.yaml"
    if not case_path.exists():
        return
    data = read_yaml(case_path)
    case_id = str(data.get("case_id") or args.case_id or "")
    previous_status = str(data.get("status") or "draft")
    rel = output.relative_to(case_dir).as_posix() if output.is_relative_to(case_dir) else str(output)
    data["application"] = {
        "status": "applied" if args.submitted else "generated",
        "artifact": rel,
        "source_template": "skills/travel-application/assets/公司出差申请单模板.docx",
        "planned_dates": f"{args.start}至{args.end}",
        "planned_route": route.split(" → "),
        "planned_purpose": args.purpose,
    }
    if args.submitted:
        data["status"] = "applied"
    docs = data.setdefault("documents", [])
    document_id = f"APP-{args.end.replace('-', '')}"
    docs[:] = [doc for doc in docs if doc.get("document_id") != document_id]
    docs.append({
        "document_id": document_id,
        "path": rel,
        "kind": "travel_application_filled",
        "source_date": args.end,
        "related_leg_ids": [],
    })
    data.setdefault("events", []).append({
        "timestamp": args.end,
        "actor": "脚本",
        "action": "生成公司出差申请单" + ("并登记为已提交" if args.submitted else ""),
        "from_status": previous_status,
        "to_status": "applied" if args.submitted else previous_status,
        "evidence": rel,
    })
    case_path.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="用本地个人资料和公司申请单模板快速生成出差申请单")
    parser.add_argument("--case-dir", help="已有差旅案件目录；提供后自动读取 case.yaml 并登记申请单")
    parser.add_argument("--output", help="输出 .docx；省略时写入案件的 01-出差申请目录")
    parser.add_argument("--case-id", default="", help="案件 ID；有 --case-dir 时默认读取 case.yaml")
    parser.add_argument("--user-profile", help="本地用户资料 YAML")
    parser.add_argument("--company-profile", help="本地公司资料 YAML，用于默认出发单位")
    parser.add_argument("--name", default="", help="出差人；默认读取本地 user-profile.yaml")
    parser.add_argument("--department", default="", help="部门；默认读取本地 user-profile.yaml")
    parser.add_argument("--title", default="", help="岗位/职别；默认读取本地 user-profile.yaml")
    parser.add_argument("--companion", default="", help="同行员工")
    parser.add_argument("--origin", default="", help="出发地；默认本地公司名称或案件 planned_route[0]")
    parser.add_argument("--destination", default="", help="目的地；默认案件 planned_route[1]")
    parser.add_argument("--purpose", default="", help="出差事由")
    parser.add_argument("--start", default="", help="计划开始日期 YYYY-MM-DD")
    parser.add_argument("--end", default="", help="计划结束日期 YYYY-MM-DD")
    parser.add_argument("--transport", default="公司车自驾", help="交通工具")
    parser.add_argument("--actual-start", default="", help="实际开始日期，可选")
    parser.add_argument("--actual-end", default="", help="实际结束日期，可选")
    parser.add_argument("--budget-lodging", dest="budget_1", default="", help="预计住宿费")
    parser.add_argument("--budget-transport", dest="budget_2", default="", help="预计交通费")
    parser.add_argument("--budget-business", dest="budget_3", default="", help="预计业务费")
    parser.add_argument("--budget-other", dest="budget_4", default="", help="预计其他备用金")
    parser.add_argument("--budget-total", dest="budget_5", default="", help="预计合计")
    parser.add_argument("--submitted", action="store_true", help="明确表示申请已经提交；会将案件状态推进到 applied")
    args = parser.parse_args()

    case_dir = Path(args.case_dir).expanduser().resolve() if args.case_dir else None
    case_data: dict[str, Any] = {}
    if case_dir and (case_dir / "case.yaml").exists():
        case_data = read_yaml(case_dir / "case.yaml")
    profile, profile_path = load_user_profile(args.user_profile)
    company_name = load_company_name(args.company_profile)

    args.case_id = args.case_id or str(case_data.get("case_id") or "TRIP-UNASSIGNED")
    args.name = args.name or profile["name"]
    args.department = args.department or profile["department"]
    args.title = args.title or profile["title"]
    planned_route = case_data.get("planned_route") or []
    args.origin = args.origin or (str(planned_route[0]) if planned_route else company_name)
    args.destination = args.destination or (str(planned_route[1]) if len(planned_route) > 1 else "")
    args.purpose = args.purpose or str(case_data.get("purpose") or "")
    args.start = args.start or str(case_data.get("planned_start") or "")
    args.end = args.end or str(case_data.get("planned_end") or "")
    args.actual_start = args.actual_start or str(case_data.get("actual_start") or "")
    args.actual_end = args.actual_end or str(case_data.get("actual_end") or "")

    missing = [name for name, value in (("--destination", args.destination), ("--purpose", args.purpose), ("--start", args.start), ("--end", args.end)) if not value]
    if missing:
        parser.error("缺少参数：" + ", ".join(missing))
    date_range(args.start, args.end)
    route = " → ".join([args.origin or "待补充", args.destination, args.origin or "待补充"])

    template = Path(__file__).resolve().parents[1] / "assets" / "公司出差申请单模板.docx"
    if not template.exists():
        raise FileNotFoundError(template)
    if args.output:
        output = Path(args.output).expanduser().resolve()
    elif case_dir:
        output = case_dir / "01-出差申请" / f"出差申请单_{args.case_id}_{args.end.replace('-', '')}.docx"
    else:
        output = Path.cwd() / f"出差申请单_{args.case_id}_{args.end.replace('-', '')}.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        raise FileExistsError(f"拒绝覆盖已有文件：{output}")

    doc = Document(template)
    filled = sum(fill_form_table(table, args, route=route, trip_days=date_range(args.start, args.end)) for table in doc.tables)
    if filled != 2:
        raise RuntimeError(f"公司模板应有两份可填写表格，实际找到 {filled} 份")
    doc.save(output)
    if case_dir:
        update_case(case_dir, output=output, args=args, route=route)
    print(output)
    print(f"user_profile={profile_path or 'none'}")
    print(f"case_updated={'yes' if case_dir else 'no'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
